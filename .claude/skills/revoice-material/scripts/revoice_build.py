"""Funnel re-voicing package builder (Steps 2, 3, 4, 5 of revoice-material-process.md).
Usage: python revoice_build.py "<Funnel dir>" [Client ...]
       python revoice_build.py --all      (desktop relay: every Funnel folder under Client Outreach)
Per client with a sources.txt: download source if missing (as "<Client> - SOURCE - <title>.mp4"),
Demucs NOVOX if missing, Whisper script .docx if no script exists, Ableton project from template
if missing, MAKE VIDEO button in Ready to Link. Idempotent: only fills gaps; complete packages are
skipped without touching the network.
Desktop relay (Spencer, 2026-10-06): YouTube blocks the cloud's IP, not this PC's. When a cloud run
can't download its pick, it uploads sources.txt (URL on line 1) to the client folder and moves on;
`--all` on the desktop finishes the package and Drive sync uploads it.
"""
import gzip, json, os, re, shutil, subprocess, sys, tempfile, zlib
import xml.etree.ElementTree as ET

FF = r"C:\Users\spenc\AppData\Local\Microsoft\WinGet\Packages\Gyan.FFmpeg_Microsoft.Winget.Source_8wekyb3d8bbwe\ffmpeg-9.0.1-full_build\bin"
DENO = r"C:\Users\spenc\AppData\Roaming\Python\Python314\Scripts"
os.environ["PATH"] = FF + os.pathsep + DENO + os.pathsep + os.environ["PATH"]
os.environ["PYTHONIOENCODING"] = "utf-8"  # titles with full-width ':' / '|' crashed cp1252 output (Klarna, Chime, 10-06)
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
# Drive for desktop moved from D: to E: (seen 2026-10-20); use whichever drive holds My Drive.
TEMPLATE = next((p for p in (r"E:\My Drive\Client Outreach\Ableton Template Project",
                             r"D:\My Drive\Client Outreach\Ableton Template Project") if os.path.isdir(p)),
                r"E:\My Drive\Client Outreach\Ableton Template Project")
WORK = os.path.join(tempfile.gettempdir(), "revoice_work")
FMT = "bestvideo[height<=1080][vcodec^=avc1]+bestaudio[ext=m4a]/bestvideo[height<=1080]+bestaudio/best[height<=1080]/best"


def log(*a):
    print(*a, flush=True)


def run(cmd, **kw):
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", **kw)
    if r.returncode != 0:
        raise RuntimeError(f"{cmd[0]} failed ({r.returncode}): {r.stderr[-800:]}")
    return r.stdout


def probe(path):
    d = json.loads(run(["ffprobe", "-v", "error", "-show_entries", "format=duration:stream=codec_type,sample_rate", "-of", "json", path]))
    dur = float(d["format"]["duration"])
    sr = next((int(s["sample_rate"]) for s in d["streams"] if s.get("codec_type") == "audio"), 44100)
    return dur, sr


def videos(cdir):
    return sorted(f for f in os.listdir(cdir) if f.lower().endswith((".mp4", ".mov", ".webm")) and "NOVOX" not in f and "no_vocals" not in f)


def read_sources(sf):
    """sources.txt: first http line = the pick. Optional lines: 'referer: <url>' (Vimeo embeds),
    'date: YYYY-MM-DD' (upload date when the host hides it, e.g. Vimeo player pages)."""
    lines = [l.strip() for l in open(sf, encoding="utf-8-sig") if l.strip()]
    url = next(l for l in lines if l.startswith("http"))
    opt = {k: l.split(":", 1)[1].strip() for l in lines for k in ("referer", "date") if l.lower().startswith(k + ":")}
    return url, opt


def ytdlp_target(url, opt):
    """Vimeo pages say 'only works when logged-in'; the player embed URL with a referer works (skill, 09-24)."""
    m = re.match(r"https?://(?:www\.)?vimeo\.com/(\d+)", url)
    if m:
        return ["--referer", opt.get("referer", "https://vimeo.com/"), "--impersonate", "chrome",
                f"https://player.vimeo.com/video/{m.group(1)}"]
    return [url]


def download(client, cdir, url, opt):
    for f in os.listdir(cdir):
        if f.endswith((".part", ".ytdl")) or ".part-Frag" in f:
            os.remove(os.path.join(cdir, f))
    # cloud_ableton.py links by this exact name: "<Client> - SOURCE - <title>.mp4"
    out = os.path.join(cdir, f"{client} - SOURCE - %(title)s.%(ext)s")
    # 403s hit some YouTube formats (avc1 SABR streams); fall back to any format, then to the default pick
    for fmt in (FMT, "bestvideo[height<=1080]+bestaudio/best[height<=1080]/best", None):
        try:
            run([sys.executable, "-m", "yt_dlp"] + (["-f", fmt] if fmt else []) +
                ["--merge-output-format", "mp4", "--windows-filenames", "--ffmpeg-location", FF, "-o", out]
                + ytdlp_target(url, opt))
            return
        except RuntimeError as e:
            err = e
            log(f"  download retry ({str(e).splitlines()[-1][:120]})")
    raise err


PREFER_DAYS = 3 * 365  # Spencer, 2026-10-05: prefer the last 3 years
MAX_AGE_DAYS = 5 * 365  # 3-5 years is OK when nothing newer fits; over 5 = rejected (NPR pick was from 2013)
# Spencer's own per-client calls to allow one video past 5 years (like NO_SAMPLE_DOMAINS in pitch_draft_qc.py).
# Only Spencer adds entries here: client folder name -> YouTube id.
AGE_EXEMPT = {"Cash App": "qSH_t9kE4es"}  # "Let's talk Cash App" (2021-06-29); Spencer 2026-10-14: "Use"


def age_gate(url, opt=None):
    """Return a reason string if the video is older than 5 years (or its date can't be read), else None."""
    import datetime
    opt = opt or {}
    try:
        if opt.get("date"):
            up = datetime.date.fromisoformat(opt["date"])
        else:
            d = run([sys.executable, "-m", "yt_dlp", "--skip-download", "--print", "upload_date"]
                    + ytdlp_target(url, opt)).strip().splitlines()[-1]
            up = datetime.datetime.strptime(d, "%Y%m%d").date()
    except Exception as e:
        return f"AGE GATE: can't read upload date ({e}); pick a video you can date"
    age = (datetime.date.today() - up).days
    if age > MAX_AGE_DAYS:
        return f"AGE GATE: uploaded {up}, older than 5 years; pick a newer video"
    if age > PREFER_DAYS:
        log(f"AGE NOTE: uploaded {up} (3-5 years old). Allowed only if nothing within 3 years fits; say so in the report.")
    return None


def novox(client, cdir, src):
    out = os.path.join(cdir, f"{client}_NOVOX.mp4")
    if os.path.exists(out):
        return out
    tmp = tempfile.mkdtemp(dir=WORK)
    wav = os.path.join(tmp, "audio.wav")
    run(["ffmpeg", "-y", "-i", src, "-vn", "-ac", "2", "-ar", "44100", wav])
    run([sys.executable, "-m", "demucs", "--two-stems=vocals", "-o", os.path.join(tmp, "stems"), wav])
    nv = os.path.join(tmp, "stems", "htdemucs", "audio", "no_vocals.wav")
    run(["ffmpeg", "-y", "-i", src, "-i", nv, "-map", "0:v:0", "-map", "1:a:0", "-c:v", "copy", "-c:a", "aac", "-shortest", out])
    shutil.move(os.path.join(tmp, "stems", "htdemucs", "audio", "vocals.wav"), vocals_path(client))  # for whisper_script
    shutil.rmtree(tmp, ignore_errors=True)
    return out


def vocals_path(client):
    return os.path.join(WORK, re.sub(r"[^A-Za-z0-9]+", "_", client) + "_vocals.wav")


def vocals(client, src):
    """Demucs vocal stem of the source (kept from the NOVOX pass; rebuilt if missing)."""
    vp = vocals_path(client)
    if not os.path.exists(vp):
        tmp = tempfile.mkdtemp(dir=WORK)
        wav = os.path.join(tmp, "audio.wav")
        run(["ffmpeg", "-y", "-i", src, "-vn", "-ac", "2", "-ar", "44100", wav])
        run([sys.executable, "-m", "demucs", "--two-stems=vocals", "-o", os.path.join(tmp, "stems"), wav])
        shutil.move(os.path.join(tmp, "stems", "htdemucs", "audio", "vocals.wav"), vp)
        shutil.rmtree(tmp, ignore_errors=True)
    return vp


def set_attr(seg, tag, value):
    # file names go into XML attributes: "&" in a title broke the .als (Qualtrics "XM Data & AI", 10-08)
    value = str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
    return re.sub(r'(<%s Value=")[^"]*(" />)' % tag, lambda m: m.group(1) + value + m.group(2), seg, count=1)


def patch_clip(xml, track_name, fpath, abs_path=None):
    """Point the one clip on `track_name` at `fpath`, written the way a manual drag-in writes it:
    clip gain 0 dB (template placeholders sit at -70 dB), unwarped, named after the file.
    Unwarped clips: CurrentEnd is in beats (120 BPM -> s*2); loop/marker fields are in seconds.
    abs_path: the path Ableton should record (e.g. Spencer's D:/My Drive/... path when the project
    is built in the cloud); defaults to fpath."""
    dur, sr = probe(fpath)
    for m in re.finditer(r"<AudioTrack Id=\"\d+\"[^>]*>", xml):
        tend = xml.find("</AudioTrack>", m.start())
        track = xml[m.start():tend]
        if f'<EffectiveName Value="{track_name}" />' not in track:
            continue
        cm = re.search(r"<AudioClip [^>]*>.*?</AudioClip>", track, re.S)
        seg = cm.group(0)
        seg = set_attr(seg, "Name", os.path.splitext(os.path.basename(fpath))[0])
        seg = set_attr(seg, "CurrentStart", 0)
        seg = set_attr(seg, "CurrentEnd", dur * 2)
        seg = set_attr(seg, "LoopStart", 0)
        seg = set_attr(seg, "StartRelative", 0)
        for t in ("LoopEnd", "OutMarker", "HiddenLoopEnd", "RightTime"):
            seg = set_attr(seg, t, dur)
        seg = set_attr(seg, "IsWarped", "false")
        seg = set_attr(seg, "SampleVolume", 1)
        seg = set_attr(seg, "RelativePath", "../" + os.path.basename(fpath))
        seg = set_attr(seg, "Path", (abs_path or fpath).replace("\\", "/"))
        seg = set_attr(seg, "OriginalFileSize", os.path.getsize(fpath))
        seg = set_attr(seg, "DefaultDuration", int(round(dur * sr)))
        seg = set_attr(seg, "DefaultSampleRate", sr)
        track = track[:cm.start()] + seg + track[cm.end():]
        return xml[:m.start()] + track + xml[tend:]
    raise RuntimeError(f"track {track_name} not found")


def read_als(als):
    """Ableton's gzip can carry trailing bytes that gzip.open rejects (Airwallex, Booking.com 10-07)."""
    return zlib.decompressobj(31).decompress(open(als, "rb").read()).decode("utf-8")


def repatch(als, src, nv):
    """Re-point an existing project's two video tracks (idempotent)."""
    xml = read_als(als)
    xml = patch_clip(xml, "ORIGINAL VIDEO", src)
    xml = patch_clip(xml, "VIDEO WITHOUT VOX", nv)
    ET.fromstring(xml.encode("utf-8"))
    with gzip.GzipFile(als, "wb", mtime=0) as g:
        g.write(xml.encode("utf-8"))
    ET.fromstring(read_als(als).encode("utf-8"))


def ableton(client, cdir, src, nv, n=1):
    pname = f"{client} Project" if n == 1 else f"{client} {n} Project"
    pdir = os.path.join(cdir, pname)
    if os.path.exists(pdir):
        return pdir
    shutil.copytree(TEMPLATE, pdir, ignore=shutil.ignore_patterns("Desktop.ini", "desktop.ini", "Backup"))
    old = os.path.join(pdir, "Ableton Template.als")
    als = os.path.join(pdir, f"{pname}.als")
    os.rename(old, als)
    xml = read_als(als)
    ET.fromstring(xml.encode("utf-8"))
    xml = patch_clip(xml, "ORIGINAL VIDEO", src)
    xml = patch_clip(xml, "VIDEO WITHOUT VOX", nv)
    ET.fromstring(xml.encode("utf-8"))
    with gzip.GzipFile(als, "wb", mtime=0) as g:
        g.write(xml.encode("utf-8"))
    ET.fromstring(read_als(als).encode("utf-8"))
    return pdir


BUTTON = r"""@echo off
title Make Video - {client}
echo.
echo  {client} : putting your newest WAV mix onto the original video...
echo.
"C:\Users\spenc\AppData\Local\Programs\Python\Python314\python.exe" "C:\Users\spenc\Bingo\.claude\skills\revoice-material\scripts\mux_take.py" "%~dp0.."
if errorlevel 1 (
  echo.
  echo  Something went wrong. Make sure you exported a .wav from Ableton into this folder.
) else (
  echo.
  echo  Done. The .mp4 is in this folder, named after your .wav.
)
echo.
pause
"""


def button(client, cdir):
    """Spencer's one-click MAKE VIDEO button in Ready to Link (muxes his WAV onto the source). Every package needs it."""
    rdir = os.path.join(cdir, "Ready to Link to Drafted Email")
    os.makedirs(rdir, exist_ok=True)
    bat = os.path.join(rdir, "MAKE VIDEO - double-click me.bat")
    if not os.path.exists(bat):
        with open(bat, "w", encoding="ascii", newline="\r\n") as f:
            f.write(BUTTON.format(client=client))
    return bat


BAD_SCRIPT = ("no captions", "transcribe manually", "transcription failed", "summary", "structure above", "context:")


def script_gate(cdir):
    """SCRIPT GATE. Every _script.docx must hold the actual VO copy of THIS folder's source video, word for word:
    no placeholders, and script_check.py (a fresh Whisper pass diffed against the script) must pass.
    Spencer, 2026-10-12, after Airbnb's script dropped its tag line "Now you can Airbnb more than an Airbnb.":
    "Make the script writing fool proof." A Google Doc script can't be read from disk, so it no longer counts as
    checked (that's how Airbnb slipped through): it reports CHECK until its text export passes --text. It is
    never converted to .docx (Spencer, 2026-10-13)."""
    import docx
    from script_check import script_text, verify
    out = []
    vids = videos(cdir)
    for sp in sorted(f for f in os.listdir(cdir) if f.endswith("_script.docx") and not f.startswith("~$")):
        paras = [p.text.strip() for p in docx.Document(os.path.join(cdir, sp)).paragraphs if p.text.strip()]
        body = " ".join(paras[1:]).lower()
        bad = [b for b in BAD_SCRIPT if b in body]
        if bad or not vids:
            out.append(f"FAIL {sp}: {'placeholder: ' + ', '.join(bad) if bad else 'no source video to check against'}")
            continue
        try:
            ok, line, gaps = verify(os.path.join(cdir, vids[0]), script_text(os.path.join(cdir, sp)))
        except Exception as e:
            ok, line, gaps = False, f"SCRIPT CHECK could not run ({e})", []
        out.append(f"{'OK' if ok else 'FAIL'} {sp}: {line}" + "".join("\n      " + g for g in gaps))
        # FORMAT: run-on lines / stray spaces (script_lines.lint). Reported, never auto-fixed on an existing
        # script: Spencer OKs any re-layout (10-13). New scripts come out of script_lines.py already clean.
        from script_lines import lint
        iss = lint([p.text for p in docx.Document(os.path.join(cdir, sp)).paragraphs])
        if iss:
            out.append(f"FORMAT {sp}: {len(iss)} issue(s), re-layout with script_lines.py only with Spencer's OK"
                       + "".join("\n      " + i for i in iss[:6]))
    if not out and has_script(cdir):
        # Never convert or rewrite it to get past this (Spencer, 2026-10-13: the .docx conversion "made the
        # scripts worse"; Funnel I was reverted to its Google Docs). Check it from a text export instead.
        return ["CHECK script is a Google Doc: export its text (FGAC drive export text/plain) and run "
                "script_check.py --src <source> --text <export>. Leave the Doc as it is; never convert it to .docx"]
    return out or ["FAIL no _script.docx"]


def has_script(cdir):
    return any(f.endswith(("_script.docx", "_script.gdoc")) and not f.startswith("~$") for f in os.listdir(cdir))


def whisper_script(client, cdir, src):
    """Write <Client>_script.docx from a local Whisper pass (desktop relay: no YouTube captions needed).
    script_lines.py lays it out as heard: one line per spoken phrase, broken at the narrator's pauses and sentence
    ends, whitespace cleaned, brand name spelled right (Spencer, 2026-10-14: "present the lines as they are heard").
    Medium model on the Demucs vocal stem, condition_on_previous_text=False, nothing past the video's end or
    Whisper's stock filler. The SCRIPT GATE then re-checks it with an independent pass (script_check.py)."""
    from script_lines import build
    build(vocals(client, src), src, os.path.join(cdir, f"{client}_script.docx"),
          header="Original copy - from Whisper transcription (desktop relay), verify against the video.",
          client=client)


def narrator(cdir, src, nv):
    """Narrator gate on a fresh download; the verdict + contact sheet go in the summary for a look."""
    sheet = os.path.join(WORK, os.path.basename(cdir) + "_frames.jpg")
    r = subprocess.run([sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)), "narrator_check.py"),
                        "--src", src, "--novox", nv, "--sheet", sheet],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    voice = next((l for l in r.stdout.splitlines() if l.startswith("VOICE")), "VOICE: ?")
    return f"{voice} | sheet {sheet}"


def archived(cdir, r):
    """True for anything under a _-prefixed subfolder (_old (...), _bad_novox, _rejected - ...): an old .als
    there must not make a package look complete (Qualtrics, 10-08) or be NOVOX-gated."""
    return any(part.startswith("_") for part in os.path.relpath(r, cdir).split(os.sep))


def complete(client, cdir):
    return bool(videos(cdir)) and os.path.exists(os.path.join(cdir, f"{client}_NOVOX.mp4")) and has_script(cdir) \
        and any(f.endswith(".als") for r, _, fs in os.walk(cdir) if not archived(cdir, r) for f in fs) \
        and os.path.exists(os.path.join(cdir, "Ready to Link to Drafted Email", "MAKE VIDEO - double-click me.bat"))


def novox_gate(client, cdir):
    """NOVOX MATCH GATE (Spencer, 2026-10-07: The Economist + BUCK NOVOX were from other videos, "this is
    killing my flow"). Every project in the folder must play this client's own source and its own NOVOX
    (novox_check.py: audio correlation + same picture). A single-video folder with a mismatch is rebuilt
    on the spot (bad file kept in _bad_novox/, never deleted); anything still failing is reported FAIL."""
    from novox_check import check_project
    projs = sorted(os.path.join(r, f) for r, _, fs in os.walk(cdir) for f in fs
                   if f.endswith(".als") and "Backup" not in r and not archived(cdir, r) and not re.search(r"\[\d{4}-", f))
    msgs, ok_all = [], True
    for als in projs:
        ok, msg = check_project(als)
        if not ok and len(videos(cdir)) == 1:
            nv = os.path.join(cdir, f"{client}_NOVOX.mp4")
            log(client, "NOVOX MISMATCH -> rebuilding:", msg)
            if os.path.exists(nv):
                bad = os.path.join(cdir, "_bad_novox"); os.makedirs(bad, exist_ok=True)
                shutil.move(nv, os.path.join(bad, os.path.basename(nv)))
                if os.path.exists(nv + ".asd"):
                    os.remove(nv + ".asd")  # Ableton's waveform cache of the wrong file
            src = os.path.join(cdir, videos(cdir)[0])
            repatch(als, src, novox(client, cdir, src))
            ok, msg = check_project(als)
            msg = "REBUILT, " + msg
        ok_all &= ok
        msgs.append(("NOVOX OK " if ok else "NOVOX MISMATCH ") + msg)
    return ok_all, msgs


def apply_swap(client, cdir):
    """REVOICE SWAP.txt (Spencer, 2026-10-07, Garmin fenix 9 vs Venu 4): the draft's hook names a different video
    than the package. Archive the old package into _old (...), turn the swap file into sources.txt, rebuild."""
    sw = os.path.join(cdir, "REVOICE SWAP.txt")
    if not os.path.exists(sw):
        return False
    lines = [l.strip() for l in open(sw, encoding="utf-8-sig") if l.strip()]
    old = next((l.split(":", 1)[1].strip() for l in lines if l.lower().startswith("replaces:")), "old pick")
    old = re.sub(r'[\\/:*?"<>|]', "", old)[:60]
    arch = os.path.join(cdir, f"_old ({old}, replaced {__import__('datetime').date.today()})")
    rdir = os.path.join(cdir, "Ready to Link to Drafted Email")
    os.makedirs(os.path.join(arch, "Ready to Link to Drafted Email"), exist_ok=True)
    for f in os.listdir(cdir):
        fp = os.path.join(cdir, f)
        if f in ("Ready to Link to Drafted Email", "REVOICE SWAP.txt", "desktop.ini") or f.startswith("_"):
            continue
        shutil.move(fp, os.path.join(arch, f))
    for f in os.listdir(rdir):
        if f.lower().endswith((".mp4", ".wav")):
            shutil.move(os.path.join(rdir, f), os.path.join(arch, "Ready to Link to Drafted Email", f))
    with open(os.path.join(cdir, "sources.txt"), "w", encoding="utf-8") as g:
        g.writelines(l + "\n" for l in lines if not l.lower().startswith("replaces:"))
    os.remove(sw)
    log(client, "REVOICE SWAP: old package archived to", os.path.basename(arch))
    return True


def build_funnel(fdir, only, results):
    for client in sorted(os.listdir(fdir)):
        cdir = os.path.join(fdir, client)
        if not os.path.isdir(cdir) or client.startswith("_") or (only and client not in only):
            continue
        key = f"{os.path.basename(fdir)}/{client}"
        apply_swap(client, cdir)
        sf = os.path.join(cdir, "sources.txt")
        if not os.path.exists(sf):
            if only:
                results[key] = "SKIP - no sources.txt (needs a pick)"
            continue  # whole-funnel / --all: no pick yet, nothing for this script to do
        if complete(client, cdir):
            ok, msgs = novox_gate(client, cdir)  # finished packages are re-checked too
            gate = [g for g in script_gate(cdir) if g.startswith("FAIL")]  # and so are their scripts
            if not ok or gate:
                results[key] = "FAIL - " + " / ".join(msgs + gate)
            continue
        try:
            url, opt = read_sources(sf)
            note = ""
            fresh = not videos(cdir)
            if fresh:
                too_old = None if AGE_EXEMPT.get(client, "~") in url else age_gate(url, opt)
                if too_old:
                    results[key] = "FAIL - " + too_old
                    log(key, "->", results[key])
                    continue
                log(key, "downloading", url)
                download(client, cdir, url, opt)
            src = os.path.join(cdir, videos(cdir)[0])
            log(key, "NOVOX from", os.path.basename(src))
            nv = novox(client, cdir, src)
            if not has_script(cdir):
                log(key, "Whisper script")
                whisper_script(client, cdir, src)
                fresh = True
            if fresh:  # the cloud couldn't hear this video, so its narrator gate never ran
                note = " | " + narrator(cdir, src, nv)
            log(key, "Ableton")
            ableton(client, cdir, src, nv)
            button(client, cdir)
            ok, msgs = novox_gate(client, cdir)
            gate = script_gate(cdir)
            ok &= not any(g.startswith("FAIL") for g in gate)
            dur, _ = probe(src)
            results[key] = (f"OK - {os.path.basename(src)} ({dur:.0f}s){note}" if ok else "FAIL") + " | " + " / ".join(msgs + gate)
        except Exception as e:
            results[key] = f"FAIL - {e}"
        log(key, "->", results[key])


def main():
    os.makedirs(WORK, exist_ok=True)
    if sys.argv[1] == "--all":
        root = os.path.dirname(TEMPLATE)
        # Funnel H onward only: A-G2 are sent or desktop-owned and use older file names
        fdirs = [os.path.join(root, d) for d in sorted(os.listdir(root)) if d.startswith("Funnel ") and d[7:8] >= "H"]
        only = set()
    else:
        fdirs, only = [sys.argv[1]], set(sys.argv[2:])
    results = {}
    for fdir in fdirs:
        build_funnel(fdir, only, results)
    log("\n=== SUMMARY ===")
    for k, v in results.items():
        log(f"{k}: {v}")
    if not results:
        log("nothing to do")
    # SCRIPT GATE: script_gate() (script_check.py on every _script.docx) runs inside build_funnel for every
    # package, finished ones included, and a miss FAILs that package's SUMMARY line above.


if __name__ == "__main__":
    main()
